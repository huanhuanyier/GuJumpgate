import tempfile
import unittest
from unittest.mock import patch
from pathlib import Path

try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None

from scripts import hotmail_helper


@unittest.skipIf(openpyxl is None, "openpyxl is not installed")
class HotmailHelperExcelTest(unittest.TestCase):
    def make_workbook(self, rows):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        for row in rows:
            sheet.append(row)
        handle = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        handle.close()
        path = Path(handle.name)
        workbook.save(path)
        return path

    def test_import_hotmail_accounts_excel_reads_column_a_raw_lines(self):
        path = self.make_workbook([
            ["mail", "name", "card", "phone", "pass", "sell"],
            ["bad----password----client----token"],
            ["user@example.com----secret----client-id----refresh-token"],
        ])

        result = hotmail_helper.import_hotmail_accounts_excel(str(path))

        self.assertEqual(result["importedCount"], 1)
        self.assertEqual(result["accounts"][0]["raw"], "user@example.com----secret----client-id----refresh-token")
        self.assertEqual(result["accounts"][0]["excelSource"]["rowNumber"], 3)

    def test_import_hotmail_accounts_excel_skips_rows_with_pass_value(self):
        path = self.make_workbook([
            ["mail", "name", "card", "phone", "pass", "sell"],
            ["done@example.com----secret----client-id----refresh-token", "", "", "", "开通", ""],
            ["failed@example.com----secret----client-id----refresh-token", "", "", "", "未开通", ""],
            ["next@example.com----secret----client-id----refresh-token", "", "", "", "", ""],
        ])

        result = hotmail_helper.import_hotmail_accounts_excel(str(path))

        self.assertEqual(result["importedCount"], 1)
        self.assertEqual(result["accounts"][0]["raw"], "next@example.com----secret----client-id----refresh-token")
        self.assertEqual(result["accounts"][0]["excelSource"]["rowNumber"], 4)

    def test_hosted_sms_excel_import_and_increment(self):
        path = self.make_workbook([
            ["phone", "link", "success"],
            ["+1 (234) 567-8901", "https://example.test/code", 2],
        ])

        result = hotmail_helper.import_hosted_sms_excel(str(path))

        self.assertEqual(result["entries"][0]["phone"], "+1 (234) 567-8901")
        self.assertEqual(result["entries"][0]["verificationUrl"], "https://example.test/code")
        self.assertEqual(result["entries"][0]["successCount"], 2)

        updated = hotmail_helper.increment_hosted_sms_excel_row(str(path), "", 2)
        workbook = openpyxl.load_workbook(path)
        self.assertEqual(updated["successCount"], 3)
        self.assertEqual(workbook.active.cell(row=2, column=3).value, 3)

    def test_update_hotmail_account_excel_row_writes_subscription_result(self):
        path = self.make_workbook([
            ["mail", "name", "card", "phone", "pass", "sell"],
            ["user@example.com----secret----client-id----refresh-token", "", "", "", "", ""],
        ])

        result = hotmail_helper.update_hotmail_account_excel_row(
            str(path),
            "",
            2,
            name="James Smith",
            card="4111111111111111",
            phone="2345678901",
            pass_status="开通",
        )

        workbook = openpyxl.load_workbook(path)
        sheet = workbook.active
        self.assertTrue(result["ok"])
        self.assertEqual(sheet.cell(row=2, column=2).value, "James Smith")
        self.assertEqual(sheet.cell(row=2, column=3).value, "4111111111111111")
        self.assertEqual(sheet.cell(row=2, column=4).value, "2345678901")
        self.assertEqual(sheet.cell(row=2, column=5).value, "开通")

    def test_unsupported_post_path_does_not_report_missing_hotmail_credentials(self):
        class Handler(hotmail_helper.HotmailHelperHandler):
            def __init__(self):
                pass

        handler = Handler()
        handler.path = "/unknown-excel-path"
        captured = {}

        with patch.object(hotmail_helper, "read_json_payload", return_value={}), \
                patch.object(hotmail_helper, "json_response", side_effect=lambda _self, status, payload: captured.update({
                    "status": status,
                    "payload": payload,
                })):
            handler.do_POST()

        self.assertEqual(captured["status"], 404)
        self.assertIn("Unsupported path", captured["payload"]["error"])


if __name__ == "__main__":
    unittest.main()
